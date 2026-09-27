import os
from pathlib import Path
from typing import Dict, Any, List, Optional
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from app.config import OUTPUT_DIR, DEFAULT_STYLE

def hex_to_rgb(hex_str: str) -> RGBColor:
    hex_str = hex_str.lstrip("#")
    if len(hex_str) != 6:
        hex_str = "1A365D"
    return RGBColor(int(hex_str[:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))

def set_cell_background(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex.lstrip("#")}"/>')
    tcPr.append(shd)

class DocumentBuilder:
    """
    Synthesizes the final output document (.docx) using python-docx.
    Applies template formatting, inserts structured text, embeds screenshots intact,
    and inserts the generated process flow diagram.
    """

    @staticmethod
    def build_document(
        structured_data: Dict[str, Any],
        style_info: Dict[str, Any],
        diagram_image_path: Optional[Path],
        screenshot_paths: Dict[str, Path],
        output_filename: str = "Formatted_Standard_Document.docx"
    ) -> Path:
        output_path = OUTPUT_DIR / output_filename
        doc = Document()

        # 1. Page Margins & Section Setup
        section = doc.sections[0]
        section.top_margin = Inches(style_info.get("margin_top_inch", 1.0))
        section.bottom_margin = Inches(style_info.get("margin_bottom_inch", 1.0))
        section.left_margin = Inches(style_info.get("margin_left_inch", 1.0))
        section.right_margin = Inches(style_info.get("margin_right_inch", 1.0))

        # Fonts & Colors
        font_primary = style_info.get("font_primary", "Calibri")
        font_heading = style_info.get("font_heading", "Calibri Light")
        primary_rgb = hex_to_rgb(style_info.get("primary_color", "#1A365D"))
        secondary_rgb = hex_to_rgb(style_info.get("secondary_color", "#2B6CB0"))
        text_rgb = hex_to_rgb(style_info.get("text_color", "#2D3748"))

        # Configure Header & Footer
        if style_info.get("header_text"):
            header = section.header
            hp = header.paragraphs[0]
            hp.text = style_info["header_text"]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if hp.runs:
                hp.runs[0].font.name = font_primary
                hp.runs[0].font.size = Pt(8.5)
                hp.runs[0].font.color.rgb = RGBColor(128, 128, 128)

        if style_info.get("footer_text"):
            footer = section.footer
            fp = footer.paragraphs[0]
            fp.text = style_info["footer_text"]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if fp.runs:
                fp.runs[0].font.name = font_primary
                fp.runs[0].font.size = Pt(8.5)
                fp.runs[0].font.color.rgb = RGBColor(128, 128, 128)

        # 2. Cover / Header Banner
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(12)
        title_p.paragraph_format.space_after = Pt(4)
        run_title = title_p.add_run(structured_data.get("title", "Standard Operating Document"))
        run_title.font.name = font_heading
        run_title.font.size = Pt(26)
        run_title.font.bold = True
        run_title.font.color.rgb = primary_rgb

        subtitle_p = doc.add_paragraph()
        subtitle_p.paragraph_format.space_after = Pt(18)
        run_sub = subtitle_p.add_run(structured_data.get("subtitle", "Synthesized Operating Standard"))
        run_sub.font.name = font_primary
        run_sub.font.size = Pt(13)
        run_sub.font.color.rgb = secondary_rgb

        # Metadata Box (Table)
        meta_table = doc.add_table(rows=1, cols=3)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = True
        hdr_cells = meta_table.rows[0].cells
        hdr_cells[0].text = f"Version: {structured_data.get('version', '1.0')}"
        hdr_cells[1].text = f"Date: {structured_data.get('date', 'Today')}"
        hdr_cells[2].text = "Status: Approved"
        for cell in hdr_cells:
            set_cell_background(cell, "F7FAFC")
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = font_primary
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(74, 85, 104)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 3. Executive Summary
        if structured_data.get("executive_summary"):
            exec_heading = doc.add_paragraph()
            exec_run = exec_heading.add_run("Executive Summary")
            exec_run.font.name = font_heading
            exec_run.font.size = Pt(16)
            exec_run.font.bold = True
            exec_run.font.color.rgb = primary_rgb
            exec_heading.paragraph_format.space_before = Pt(14)
            exec_heading.paragraph_format.space_after = Pt(6)

            # Executive summary callout table
            summary_table = doc.add_table(rows=1, cols=1)
            summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = summary_table.rows[0].cells[0]
            set_cell_background(cell, "EBF8FF")
            p = cell.paragraphs[0]
            r = p.add_run(structured_data["executive_summary"])
            r.font.name = font_primary
            r.font.size = Pt(10.5)
            r.font.italic = True
            r.font.color.rgb = text_rgb
            doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 4. Main Document Sections
        for section_item in structured_data.get("sections", []):
            h_p = doc.add_paragraph()
            h_p.paragraph_format.space_before = Pt(16)
            h_p.paragraph_format.space_after = Pt(6)
            h_run = h_p.add_run(section_item.get("heading", "Section"))
            h_run.font.name = font_heading
            h_run.font.size = Pt(15)
            h_run.font.bold = True
            h_run.font.color.rgb = primary_rgb

            # Paragraphs
            for paragraph_text in section_item.get("paragraphs", []):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(paragraph_text)
                r.font.name = font_primary
                r.font.size = Pt(11)
                r.font.color.rgb = text_rgb

            # Bullet points
            for bullet in section_item.get("bullet_points", []):
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_after = Pt(3)
                r = bp.add_run(bullet)
                r.font.name = font_primary
                r.font.size = Pt(10.5)
                r.font.color.rgb = text_rgb

            # Callout box if present
            if section_item.get("callout_box"):
                c_table = doc.add_table(rows=1, cols=1)
                c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                c_cell = c_table.rows[0].cells[0]
                set_cell_background(c_cell, "FEFCBF")  # Light amber shading
                cp = c_cell.paragraphs[0]
                cr = cp.add_run(f"Note: {section_item['callout_box']}")
                cr.font.name = font_primary
                cr.font.size = Pt(10)
                cr.font.bold = True
                cr.font.color.rgb = RGBColor(116, 66, 16)
                doc.add_paragraph().paragraph_format.space_after = Pt(8)

            # Insert Referenced Screenshots Intact (Unmodified)
            for ref_name in section_item.get("screenshot_references", []):
                if ref_name in screenshot_paths and screenshot_paths[ref_name].exists():
                    img_p = doc.add_paragraph()
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.paragraph_format.space_before = Pt(8)
                    img_p.paragraph_format.space_after = Pt(4)
                    
                    # Add image intact
                    img_p.add_run().add_picture(str(screenshot_paths[ref_name]), width=Inches(5.5))
                    
                    # Image caption
                    cap_p = doc.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_after = Pt(12)
                    cap_r = cap_p.add_run(f"Figure: Preserved Input Image ({ref_name})")
                    cap_r.font.name = font_primary
                    cap_r.font.size = Pt(9)
                    cap_r.font.italic = True
                    cap_r.font.color.rgb = RGBColor(113, 128, 150)

        # 5. Process Flow Diagram Section
        diag_heading = doc.add_paragraph()
        diag_heading.paragraph_format.space_before = Pt(20)
        diag_heading.paragraph_format.space_after = Pt(8)
        diag_run = diag_heading.add_run(f"Process Flow Diagram: {structured_data.get('process_title', 'Standard Process')}")
        diag_run.font.name = font_heading
        diag_run.font.size = Pt(16)
        diag_run.font.bold = True
        diag_run.font.color.rgb = primary_rgb

        if diagram_image_path and diagram_image_path.exists():
            diag_p = doc.add_paragraph()
            diag_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            diag_p.paragraph_format.space_before = Pt(8)
            diag_p.paragraph_format.space_after = Pt(4)
            diag_p.add_run().add_picture(str(diagram_image_path), width=Inches(5.8))
            
            d_cap = doc.add_paragraph()
            d_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            d_cap.paragraph_format.space_after = Pt(16)
            d_run = d_cap.add_run("Figure 1.0: Generated Process Architecture Diagram")
            d_run.font.name = font_primary
            d_run.font.size = Pt(9.5)
            d_run.font.italic = True
            d_run.font.color.rgb = RGBColor(113, 128, 150)

        # Process Steps Details Table
        if structured_data.get("process_steps"):
            p_table = doc.add_table(rows=1, cols=4)
            p_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            p_table.autofit = True
            
            headers = ["Step #", "Role / Actor", "Action / Milestone", "Detailed Description"]
            hdr_row = p_table.rows[0].cells
            for idx, text in enumerate(headers):
                hdr_row[idx].text = text
                set_cell_background(hdr_row[idx], style_info.get("primary_color", "1A365D").lstrip("#"))
                for p in hdr_row[idx].paragraphs:
                    for r in p.runs:
                        r.font.name = font_primary
                        r.font.size = Pt(10)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)

            for step in structured_data.get("process_steps", []):
                row_cells = p_table.add_row().cells
                row_cells[0].text = str(step.get("step_number", ""))
                row_cells[1].text = str(step.get("actor", ""))
                row_cells[2].text = str(step.get("action_title", ""))
                row_cells[3].text = str(step.get("description", ""))
                
                for idx, cell in enumerate(row_cells):
                    set_cell_background(cell, "F7FAFC" if step.get("step_number", 0) % 2 == 0 else "FFFFFF")
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.name = font_primary
                            r.font.size = Pt(9.5)
                            r.font.color.rgb = text_rgb

        # Unreferenced Screenshots appended at the end intact
        all_referenced = []
        for s in structured_data.get("sections", []):
            all_referenced.extend(s.get("screenshot_references", []))

        unref_keys = [k for k in screenshot_paths.keys() if k not in all_referenced]
        if unref_keys:
            app_h = doc.add_paragraph()
            app_h.paragraph_format.space_before = Pt(20)
            app_h.paragraph_format.space_after = Pt(8)
            app_r = app_h.add_run("Appendix: Additional Process Screenshots")
            app_r.font.name = font_heading
            app_r.font.size = Pt(15)
            app_r.font.bold = True
            app_r.font.color.rgb = primary_rgb

            for key in unref_keys:
                if screenshot_paths[key].exists():
                    img_p = doc.add_paragraph()
                    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    img_p.paragraph_format.space_before = Pt(8)
                    img_p.add_run().add_picture(str(screenshot_paths[key]), width=Inches(5.5))
                    
                    cap_p = doc.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_after = Pt(12)
                    cap_r = cap_p.add_run(f"Input Attachment: {key}")
                    cap_r.font.name = font_primary
                    cap_r.font.size = Pt(9)
                    cap_r.font.italic = True

        doc.save(output_path)
        print(f"[DocumentBuilder] Document successfully generated at {output_path}")
        return output_path
